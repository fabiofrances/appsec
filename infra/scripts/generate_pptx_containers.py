#!/usr/bin/env python3
"""Gera o deck da palestra "Containerização na Prática" (hands-on).

Uso: python infra/scripts/generate_pptx_containers.py
Saída: reports/containerizacao-palestra.pptx
Requer: pip install python-pptx

Os números do comparativo ficam em MEASURED. Depois de rodar `make lab-compare`,
substitua os valores e mude ESTIMATED para False (ver hands-on/VALIDACAO.md).
"""

import os

from pptx import Presentation
from pptx.chart.data import CategoryChartData
from pptx.dml.color import RGBColor
from pptx.enum.chart import XL_CHART_TYPE, XL_LABEL_POSITION
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import MSO_ANCHOR, PP_ALIGN
from pptx.util import Inches, Pt

# ── Números medidos em 29/10/2026 (make lab-compare / lab-duel no WSL2) ─────
# CVEs mudam toda semana: remeça na véspera (ver hands-on/VALIDACAO.md).
ESTIMATED = False
MEASURED = {
    "backend:single": 2070,
    "backend:full": 1715,
    "api-go:single": 1375,
    "backend:slim": 477,
    "backend:chainguard": 410,
    "backend:alpine": 394,
    "backend:distroless": 361,
    "api-go:alpine": 21,
    "api-go:distroless": 15,
    "api-go:scratch": 9,
}

# Duelo Alpine x Distroless (make lab-duel)
DUEL = {
    "Tamanho": ("394 MB", "361 MB"),
    "Arquivos no filesystem": ("18.946", "17.676"),
    "Executáveis em /bin": ("314", "0"),
    "Shell / pacotes": ("sh + apk", "nenhum"),
    "libc": ("musl", "glibc"),
    "CVEs CRITICAL / HIGH": ("1 / 35 (0 no SO)", "2 / 31"),
    "docker exec sh": ("funciona", "falha"),
}

# ── Paleta ───────────────────────────────────────────────────────────────────
BG = RGBColor(0x0B, 0x1E, 0x2D)      # navy profundo
CARD = RGBColor(0x13, 0x31, 0x4A)    # "container"
CODE_BG = RGBColor(0x06, 0x12, 0x1C)
DOCKER = RGBColor(0x24, 0x96, 0xED)  # acento principal
TEAL = RGBColor(0x1F, 0xC8, 0xB4)    # "bom"
AMBER = RGBColor(0xFF, 0xB0, 0x20)   # "atenção"
CORAL = RGBColor(0xFF, 0x5A, 0x5F)   # "ruim"
WHITE = RGBColor(0xFF, 0xFF, 0xFF)
MUTED = RGBColor(0xA9, 0xBC, 0xD0)

FONT = "Calibri"
MONO = "Courier New"

prs = Presentation()
prs.slide_width = Inches(13.333)
prs.slide_height = Inches(7.5)
BLANK = prs.slide_layouts[6]


# ── Helpers ──────────────────────────────────────────────────────────────────
def new_slide(notes=""):
    s = prs.slides.add_slide(BLANK)
    s.background.fill.solid()
    s.background.fill.fore_color.rgb = BG
    if notes:
        s.notes_slide.notes_text_frame.text = notes
    return s


def text(s, x, y, w, h, paras, size=16, color=WHITE, bold=False, font=FONT,
         align=PP_ALIGN.LEFT, anchor=MSO_ANCHOR.TOP, space_after=4, margin=0.0):
    """paras: str | list[str | list[(texto, {opções})]]"""
    tb = s.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    tf = tb.text_frame
    tf.word_wrap = True
    tf.vertical_anchor = anchor
    for side in ("left", "right", "top", "bottom"):
        setattr(tf, f"margin_{side}", Inches(margin))
    if isinstance(paras, str):
        paras = [paras]
    for i, para in enumerate(paras):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.alignment = align
        p.space_after = Pt(space_after)
        runs = para if isinstance(para, list) else [(para, {})]
        for content, opt in runs:
            r = p.add_run()
            r.text = content
            f = r.font
            f.name = opt.get("font", font)
            f.size = Pt(opt.get("size", size))
            f.bold = opt.get("bold", bold)
            f.italic = opt.get("italic", False)
            f.color.rgb = opt.get("color", color)
    return tb


def rect(s, x, y, w, h, fill, rounded=True):
    shape = MSO_SHAPE.ROUNDED_RECTANGLE if rounded else MSO_SHAPE.RECTANGLE
    sh = s.shapes.add_shape(shape, Inches(x), Inches(y), Inches(w), Inches(h))
    sh.fill.solid()
    sh.fill.fore_color.rgb = fill
    sh.line.fill.background()
    sh.shadow.inherit = False
    if rounded:
        sh.adjustments[0] = 0.08
    return sh


def badge(s, x, y, label, color, d=0.5, size=16):
    c = s.shapes.add_shape(MSO_SHAPE.OVAL, Inches(x), Inches(y), Inches(d), Inches(d))
    c.fill.solid()
    c.fill.fore_color.rgb = color
    c.line.fill.background()
    c.shadow.inherit = False
    tf = c.text_frame
    for side in ("left", "right", "top", "bottom"):
        setattr(tf, f"margin_{side}", 0)
    tf.vertical_anchor = MSO_ANCHOR.MIDDLE
    p = tf.paragraphs[0]
    p.alignment = PP_ALIGN.CENTER
    r = p.add_run()
    r.text = label
    r.font.name = FONT
    r.font.size = Pt(size)
    r.font.bold = True
    r.font.color.rgb = BG


def title(s, t, kicker=None):
    if kicker:
        text(s, 0.6, 0.35, 12, 0.35, kicker.upper(), size=13, color=DOCKER, bold=True)
    text(s, 0.6, 0.62, 12.1, 0.8, t, size=36, bold=True)


def card(s, x, y, w, h, head, lines, color=DOCKER, mark=None, size=16):
    rect(s, x, y, w, h, CARD)
    tx = x + 0.25
    if mark:
        badge(s, x + 0.25, y + 0.22, mark, color, d=0.45, size=14)
        tx = x + 0.85
    text(s, tx, y + 0.22, w - (tx - x) - 0.2, 0.45, head, size=18, bold=True,
         anchor=MSO_ANCHOR.MIDDLE)
    if lines:
        text(s, x + 0.25, y + 0.8, w - 0.5, h - 0.95, lines, size=size, color=MUTED,
             space_after=5)


def code(s, x, y, w, h, lines, size=13):
    rect(s, x, y, w, h, CODE_BG)
    paras = []
    for line in lines:
        if line.startswith("#"):
            paras.append([(line, {"color": MUTED, "font": MONO})])
        else:
            paras.append([(line, {"font": MONO})])
    text(s, x + 0.25, y + 0.2, w - 0.5, h - 0.4, paras, size=size, space_after=2)


def lab(s, cmd, minutes=None):
    label = f"LAB  ·  {cmd}" + (f"  ·  {minutes} min" if minutes else "")
    w = 0.11 * len(label) + 0.6
    rect(s, 0.6, 6.7, w, 0.45, DOCKER)
    text(s, 0.6, 6.7, w, 0.45, label, size=14, bold=True, color=BG,
         align=PP_ALIGN.CENTER, anchor=MSO_ANCHOR.MIDDLE, font=MONO)


def table(s, x, y, col_w, rows, size=13, row_h=0.42, head_color=DOCKER, highlight=None):
    shape = s.shapes.add_table(len(rows), len(col_w), Inches(x), Inches(y),
                               Inches(sum(col_w)), Inches(row_h * len(rows)))
    tbl = shape.table
    for j, cw in enumerate(col_w):
        tbl.columns[j].width = Inches(cw)
    for i, row in enumerate(rows):
        tbl.rows[i].height = Inches(row_h)
        for j, val in enumerate(row):
            cell = tbl.cell(i, j)
            cell.fill.solid()
            if i == 0:
                cell.fill.fore_color.rgb = head_color
            elif highlight and i in highlight:
                cell.fill.fore_color.rgb = RGBColor(0x1B, 0x45, 0x63)
            else:
                cell.fill.fore_color.rgb = CARD if i % 2 else RGBColor(0x0F, 0x28, 0x3D)
            cell.margin_left = cell.margin_right = Inches(0.1)
            cell.margin_top = cell.margin_bottom = Inches(0.04)
            cell.vertical_anchor = MSO_ANCHOR.MIDDLE
            tf = cell.text_frame
            tf.word_wrap = True
            p = tf.paragraphs[0]
            r = p.add_run()
            r.text = val
            r.font.name = FONT
            r.font.size = Pt(size)
            r.font.bold = i == 0 or j == 0
            r.font.color.rgb = BG if i == 0 else (WHITE if j == 0 else MUTED)
    return tbl


# ═════════════════════════════════════════════════════════════════════════════
# 1 — Capa
# ═════════════════════════════════════════════════════════════════════════════
s = new_slide("Abertura. Pergunta para a sala: quem já escreveu um Dockerfile que funciona? "
              "Hoje a pergunta é outra: ele é pequeno, rápido de construir e seguro? Vamos medir ao vivo.")
for i, (cx, cy) in enumerate([(8.2, 1.3), (9.9, 1.3), (11.6, 1.3), (8.2, 2.9), (9.9, 2.9),
                              (8.2, 4.5)]):
    rect(s, cx, cy, 1.5, 1.4, [DOCKER, CARD, CARD, TEAL, CARD, CARD][i])
text(s, 0.8, 1.6, 7.2, 0.4, "HANDS-ON  ·  ~2 HORAS", size=16, bold=True, color=DOCKER)
text(s, 0.8, 2.1, 7.2, 1.9, ["Containerização", "na Prática"], size=54, bold=True, space_after=0)
text(s, 0.8, 4.15, 7.0, 1.0,
     "Do Dockerfile ingênuo à imagem mínima, escaneada, endurecida e rodando em Swarm",
     size=20, color=MUTED)
text(s, 0.8, 5.6, 7.0, 0.4, "Calculadora de IMC  ·  NestJS  ·  Go  ·  Vue + nginx",
     size=15, color=TEAL, bold=True)

# ═════════════════════════════════════════════════════════════════════════════
# 1b — Votação de abertura
# ═════════════════════════════════════════════════════════════════════════════
s = new_slide("Antes de qualquer conteúdo: votação com a sala. Anote o placar no quadro e peça um "
              "argumento de cada lado, sem responder. Prometa: no bloco 2 medimos tudo com a mesma app e, "
              "no final, votamos de novo.")
text(s, 0.6, 0.9, 12.1, 0.5, "VOTAÇÃO DE ABERTURA", size=16, bold=True, color=DOCKER,
     align=PP_ALIGN.CENTER)
text(s, 0.6, 1.45, 12.1, 1.0, "Sua API Node vai para produção amanhã.", size=34, bold=True,
     align=PP_ALIGN.CENTER)
text(s, 0.6, 2.3, 12.1, 0.6, "Qual imagem base você escolhe?", size=24, color=MUTED,
     align=PP_ALIGN.CENTER)
for i, (name, tag, args, color) in enumerate([
    ("Alpine", "node:22-alpine", ["\"é menor\"", "\"tem shell para debug\"", "\"todo mundo usa\""], AMBER),
    ("Distroless", "distroless/nodejs22", ["\"é mais seguro\"", "\"sem shell, sem ataque\"", "\"impossível debugar?\""], TEAL),
]):
    x = 1.4 + i * 5.6
    rect(s, x, 3.3, 4.95, 3.3, CARD)
    text(s, x, 3.45, 4.95, 0.7, name, size=34, bold=True, color=color, align=PP_ALIGN.CENTER)
    text(s, x, 4.15, 4.95, 0.4, tag, size=15, color=MUTED, font=MONO, align=PP_ALIGN.CENTER)
    text(s, x + 0.4, 4.75, 4.15, 1.7, args, size=17, align=PP_ALIGN.CENTER, space_after=6)
text(s, 6.35, 4.45, 0.6, 0.6, "x", size=32, bold=True, color=MUTED, align=PP_ALIGN.CENTER)

# ═════════════════════════════════════════════════════════════════════════════
# 2 — Agenda
# ═════════════════════════════════════════════════════════════════════════════
s = new_slide("Mostre o mapa da aula. Cada bloco termina com um lab. "
              "Peça para quem não rodou make lab-prepull rodar agora.")
title(s, "O caminho de hoje", "Agenda")
agenda = [
    ("0", "Abertura: Alpine x Distroless?", "8"), ("1", "Fundamentos e registries", "7"),
    ("2", "Imagens base + duelo", "15"), ("3", "Single x multi-stage", "15"),
    ("4", "Seu Dockerfile + BuildKit", "15"), ("~", "Intervalo", "5"),
    ("5", "Scan e supply chain", "15"), ("6", "Hardening do runtime", "10"),
    ("7", "Comandos de rotina", "10"), ("8", "Docker Swarm", "15"),
    ("9", "Encerramento e Q&A", "5"),
]
for i, (n, label, mins) in enumerate(agenda):
    col, row = divmod(i, 6)
    x, y = 0.6 + col * 6.2, 1.75 + row * 0.82
    rect(s, x, y, 5.9, 0.66, CARD)
    badge(s, x + 0.12, y + 0.1, n, AMBER if n == "~" else DOCKER, d=0.46, size=15)
    text(s, x + 0.75, y, 4.0, 0.66, label, size=17, bold=True, anchor=MSO_ANCHOR.MIDDLE)
    text(s, x + 4.7, y, 1.05, 0.66, f"{mins} min", size=15, color=MUTED,
         align=PP_ALIGN.RIGHT, anchor=MSO_ANCHOR.MIDDLE)
text(s, 6.8, 6.7, 5.9, 0.5, "Material: hands-on/README.md  ·  CHEATSHEET.md", size=14,
     color=MUTED, anchor=MSO_ANCHOR.MIDDLE)

# ═════════════════════════════════════════════════════════════════════════════
# 3 — O placar
# ═════════════════════════════════════════════════════════════════════════════
s = new_slide("Fio condutor: a mesma API, do pior Dockerfile possível até a imagem mínima. "
              "Esse é o placar que vamos construir ao longo da aula com make lab-compare.")
title(s, "Mesma API, duas imagens", "O placar da aula")
big = [
    ("backend:single", f"{MEASURED['backend:single']:,} MB".replace(",", "."), CORAL,
     ["node:22 full + devDeps", "roda como root", "25 CRITICAL · 537 HIGH"]),
    ("api-go:scratch", f"{MEASURED['api-go:scratch']} MB", TEAL,
     ["binário estático", "usuário 65532", "0 CVEs (Go 1.27)"]),
]
for i, (name, val, color, facts) in enumerate(big):
    x = 0.6 + i * 6.25
    rect(s, x, 1.75, 5.85, 4.6, CARD)
    text(s, x + 0.4, 1.95, 5.0, 0.45, name, size=18, bold=True, color=MUTED, font=MONO)
    text(s, x + 0.4, 2.45, 5.2, 1.4, val, size=72, bold=True, color=color)
    text(s, x + 0.4, 4.1, 5.0, 2.0, facts, size=18, color=WHITE, space_after=6)
text(s, 5.85, 3.6, 1.6, 0.8, "→", size=48, bold=True, color=AMBER, align=PP_ALIGN.CENTER)
if ESTIMATED:
    text(s, 0.6, 6.55, 12, 0.4, "≈ valores estimados — atualizar com make lab-compare",
         size=12, color=MUTED)

# ═════════════════════════════════════════════════════════════════════════════
# 4 — Fundamentos
# ═════════════════════════════════════════════════════════════════════════════
s = new_slide("Imagem = pilha de camadas read-only + config. Container = imagem + camada gravável. "
              "Apagar arquivo numa camada posterior NÃO reduz a imagem. Mostrar docker history.")
title(s, "Imagem é uma pilha de camadas", "Bloco 1 · Fundamentos")
layers = [("camada gravável do container", AMBER), ("COPY dist/  (sua app)", TEAL),
          ("RUN npm ci --omit=dev", DOCKER), ("WORKDIR /app", CARD),
          ("FROM node:22-alpine  (base)", CARD)]
for i, (label, color) in enumerate(layers):
    y = 1.8 + i * 0.85
    rect(s, 0.9, y, 5.6, 0.7, color)
    text(s, 1.1, y, 5.2, 0.7, label, size=16, bold=True, font=MONO,
         color=BG if color in (AMBER, TEAL, DOCKER) else WHITE, anchor=MSO_ANCHOR.MIDDLE)
text(s, 0.9, 6.15, 5.6, 0.4, "read-only, compartilhadas, endereçadas por hash",
     size=14, color=MUTED, align=PP_ALIGN.CENTER)
facts = [
    [("OCI  ", {"bold": True, "color": DOCKER}),
     ("a mesma imagem roda em Docker, containerd, Podman e Kubernetes", {})],
    [("Cache  ", {"bold": True, "color": DOCKER}),
     ("cada instrução gera uma camada; a ordem define o reaproveitamento", {})],
    [("Armadilha  ", {"bold": True, "color": CORAL}),
     ("rm numa camada posterior esconde o arquivo, mas ele continua na imagem", {})],
]
text(s, 7.1, 1.85, 5.6, 3.0, facts, size=17, space_after=14)
code(s, 7.1, 4.75, 5.6, 1.4, ["docker history node:22-alpine",
                              "docker image inspect node:22-alpine \\",
                              "  -f '{{json .RootFS.Layers}}'"], size=13)

# ═════════════════════════════════════════════════════════════════════════════
# 5 — Onde encontrar imagens
# ═════════════════════════════════════════════════════════════════════════════
s = new_slide("Docker Hub Official/Verified primeiro; registries de fabricante; imagens mínimas "
              "(distroless, Chainguard, Docker Hardened Images). Em produção, sempre via registry "
              "interno com scan e política.")
title(s, "Onde encontrar imagens", "Bloco 1 · Registries")
regs = [
    ("Docker Hub", ["Official Images e Verified Publisher", "atenção ao rate limit anônimo"], DOCKER),
    ("GHCR · Quay · ECR Public", ["projetos GitHub, Red Hat", "ECR Public espelha as oficiais"], DOCKER),
    ("Distroless", ["gcr.io/distroless", "só runtime, sem shell"], TEAL),
    ("Chainguard · DHI", ["cgr.dev/chainguard, dhi.io", "rebuild diário, ~0 CVEs"], TEAL),
    ("Registry interno", ["Harbor, Nexus, ECR, ACR, OCIR", "espelho + scan + política"], AMBER),
    ("Como escolher", ["oficial > fabricante > usuário", "rebuild recente, SBOM, assinatura"], AMBER),
]
for i, (head, lines, color) in enumerate(regs):
    col, row = i % 3, i // 3
    card(s, 0.6 + col * 4.1, 1.75 + row * 2.45, 3.85, 2.2, head, lines, color=color,
         mark=str(i + 1))

# ═════════════════════════════════════════════════════════════════════════════
# 6 — Tag x digest
# ═════════════════════════════════════════════════════════════════════════════
s = new_slide("Tag é um ponteiro mutável; digest é o hash do conteúdo. latest quebra a "
              "reprodutibilidade e o rollback. Renovate/Dependabot atualizam o digest via PR.")
title(s, "Tag muda. Digest não.", "Bloco 1 · Versionamento")
code(s, 0.6, 1.8, 7.3, 2.6, [
    "# mutável: muda a cada rebuild do upstream",
    "FROM node:22-alpine",
    "",
    "# imutável: sempre os mesmos bytes",
    "FROM node:22-alpine@sha256:9f3a...c1e2",
], size=15)
code(s, 0.6, 4.65, 7.3, 1.25, ["docker buildx imagetools inspect node:22-alpine"], size=15)
card(s, 8.3, 1.8, 4.4, 1.9, "Evite :latest", ["build não reprodutível", "rollback impossível"],
     color=CORAL, mark="✕")
card(s, 8.3, 4.0, 4.4, 1.9, "Pin + automação", ["digest no FROM", "Renovate abre o PR de update"],
     color=TEAL, mark="✓")

# ═════════════════════════════════════════════════════════════════════════════
# 7 — Tipos de imagem base
# ═════════════════════════════════════════════════════════════════════════════
s = new_slide("A mesma API, o mesmo código, só muda o FROM. Full para build; slim/alpine/"
              "distroless/chainguard para runtime; scratch para binários estáticos.")
title(s, "Tipos de imagem base", "Bloco 2 · 15 min")
table(s, 0.6, 1.75, [2.6, 2.6, 1.3, 2.7, 2.9], [
    ["Variante", "Base", "libc", "Shell / pacotes", "Uso típico"],
    ["node:22", "Debian + toolchain", "glibc", "bash, apt, gcc, git", "stage de build"],
    ["node:22-slim", "Debian mínimo", "glibc", "bash, apt", "runtime com glibc"],
    ["node:22-alpine", "Alpine", "musl", "busybox sh, apk", "runtime pequeno"],
    ["distroless", "Debian sem userland", "glibc", "nenhum", "runtime de produção"],
    ["chainguard", "Wolfi", "glibc", "busybox sh (:latest)", "runtime 0 CVEs de SO"],
    ["scratch", "vazio", "—", "nenhum", "binário estático"],
], size=15, row_h=0.6, highlight={3})
lab(s, "make lab-bases lab-go lab-compare")

# ═════════════════════════════════════════════════════════════════════════════
# 8 — Alpine
# ═════════════════════════════════════════════════════════════════════════════
s = new_slide("BusyBox: um binário de ~1 MB implementa sh, ls, wget via symlinks. musl: binários "
              "glibc não carregam, DNS diferente, alocador diferente. Regra: builder e runtime com "
              "a mesma libc.")
title(s, "Alpine não é \"só menor\"", "Bloco 2 · busybox + musl")
card(s, 0.6, 1.75, 6.0, 2.3, "BusyBox", [
    "1 binário (~1 MB) faz sh, ls, wget, ps...",
    "symlinks em /bin → /bin/busybox",
    "flags e saídas diferentes do GNU coreutils",
], color=DOCKER, mark="1")
card(s, 6.85, 1.75, 5.9, 2.3, "musl libc", [
    "binário glibc: \"not found\" / \"Error relocating\"",
    "módulos nativos (npm, pip) recompilam",
    "resolver DNS e alocador diferentes",
], color=AMBER, mark="2")
code(s, 0.6, 4.35, 12.15, 1.6, [
    "docker run --rm alpine readlink -f /bin/ls            # /bin/busybox",
    "docker run --rm node:22-alpine ldd /usr/local/bin/node  # musl",
    "docker run --rm node:22-slim   ldd /usr/local/bin/node  # glibc",
], size=16)
text(s, 0.6, 6.25, 12, 0.5, [[("Regra: ", {"bold": True, "color": TEAL}),
                              ("builder e runtime com a mesma libc", {})]], size=18)

# ═════════════════════════════════════════════════════════════════════════════
# 9 — Distroless / Chainguard / scratch
# ═════════════════════════════════════════════════════════════════════════════
s = new_slide("Sem shell = sem docker exec sh. Debug com docker debug ou imagem :debug. "
              "HEALTHCHECK vai para o orquestrador ou para o próprio binário (flag -healthcheck da API Go).")
title(s, "Sem shell: o que muda", "Bloco 2 · distroless · scratch")
card(s, 0.6, 1.75, 3.9, 2.55, "Debug", ["sem docker exec sh", "docker debug <ctr>", "tag :debug (busybox)"],
     color=DOCKER, mark="1")
card(s, 4.72, 1.75, 3.9, 2.55, "Healthcheck", ["sem wget/curl", "probe no orquestrador",
                                               "ou o binário se testa"], color=AMBER, mark="2")
card(s, 8.84, 1.75, 3.9, 2.55, "Usuário", ["nonroot (65532) padrão", "USER numérico no scratch",
                                          "sem /etc/passwd"], color=TEAL, mark="3")
code(s, 0.6, 4.6, 12.15, 1.25, [
    "HEALTHCHECK CMD [\"/bmi-api\", \"-healthcheck\"]    # hands-on/api-go/Dockerfile.scratch",
    "docker run --rm -it --entrypoint sh handson/backend:distroless   # falha: sem shell",
], size=14)

# ═════════════════════════════════════════════════════════════════════════════
# 9b — Duelo Alpine x Distroless
# ═════════════════════════════════════════════════════════════════════════════
s = new_slide("Resposta à votação de abertura, com dados medidos por make lab-duel. Com Node, tamanho "
              "não decide: a decisão é libc (musl e Node experimental x glibc), superfície de ataque "
              "(shell e apk disponíveis ao atacante) e capacidade de debug. Chainguard: 0 CVEs de SO, mas a "
              ":latest gratuita traz busybox/sh. "
              "Para Go, o duelo nem existe: scratch.")
title(s, "O duelo: Alpine x Distroless", "Bloco 2 · resposta à votação")
rows = [["", "Alpine", "Distroless"]] + [[k, a, b] for k, (a, b) in DUEL.items()]
table(s, 0.6, 1.75, [3.2, 2.6, 2.6], rows, size=15, row_h=0.56)
verdict = [
    [("Tamanho não decide ", {"bold": True, "color": AMBER}), ("com Node: node_modules domina", {})],
    [("Decide: ", {"bold": True, "color": TEAL}), ("libc, superfície de ataque e debug", {})],
    [("CVEs: ", {"bold": True, "color": CORAL}), ("a maioria é da sua app (node_modules)", {})],
    [("Chainguard: ", {"bold": True, "color": DOCKER}), ("0 CVEs de SO, mas a :latest traz sh", {})],
]
text(s, 9.3, 1.85, 3.45, 4.4, verdict, size=16, space_after=14)
if ESTIMATED:
    text(s, 0.6, 6.05, 8.4, 0.4, "≈ valores estimados — atualizar com make lab-duel",
         size=12, color=MUTED)
lab(s, "make lab-duel")

# ═════════════════════════════════════════════════════════════════════════════
# 10 — Gráfico comparativo
# ═════════════════════════════════════════════════════════════════════════════
s = new_slide("Projete a tabela do make lab-compare. O node_modules domina o tamanho nas imagens "
              "Node, por isso alpine/distroless/chainguard ficam parecidos. A versão Go em scratch é "
              "uma ordem de grandeza menor.")
title(s, "Tamanho da imagem final (MB)", "Bloco 2 · resultado do lab")
cd = CategoryChartData()
names = list(MEASURED)[::-1]
cd.categories = names
cd.add_series("MB", [MEASURED[n] for n in names])
gf = s.shapes.add_chart(XL_CHART_TYPE.BAR_CLUSTERED, Inches(0.6), Inches(1.6),
                        Inches(12.1), Inches(4.95), cd)
ch = gf.chart
ch.has_legend = False
ch.has_title = False
ch.font.size = Pt(13)
ch.font.name = FONT
ch.font.color.rgb = MUTED
plot = ch.plots[0]
plot.gap_width = 45
plot.has_data_labels = True
dl = plot.data_labels
dl.number_format = "#,##0"
dl.number_format_is_linked = False
dl.position = XL_LABEL_POSITION.OUTSIDE_END
dl.font.size = Pt(13)
dl.font.bold = True
dl.font.color.rgb = WHITE
for idx, n in enumerate(names):
    pt = plot.series[0].points[idx]
    pt.format.fill.solid()
    v = MEASURED[n]
    pt.format.fill.fore_color.rgb = CORAL if v >= 800 else (AMBER if v >= 100 else TEAL)
va = ch.value_axis
va.has_major_gridlines = True
va.major_gridlines.format.line.color.rgb = CARD
va.format.line.fill.background()
va.tick_labels.font.color.rgb = MUTED
ca = ch.category_axis
ca.format.line.color.rgb = CARD
ca.tick_labels.font.color.rgb = WHITE
ca.tick_labels.font.name = MONO
if ESTIMATED:
    text(s, 0.6, 6.65, 12, 0.4, "≈ valores estimados — atualizar com make lab-compare",
         size=12, color=MUTED)

# ═════════════════════════════════════════════════════════════════════════════
# 10b — De onde vêm as CVEs
# ═════════════════════════════════════════════════════════════════════════════
s = new_slide("O argumento que mais surpreende: nas bases mínimas, a maior parte das CVEs é da app "
              "(node_modules; o CRITICAL é proxy-addr, dependência do Express). A node:22-alpine ainda traz "
              "o npm, que soma ~10 HIGH. E a história do Go: com 1.24 (fora de suporte), 23 HIGH até no "
              "scratch, vindos da stdlib embutida; com 1.27, zero.")
title(s, "As CVEs são suas", "Bloco 2 · de onde vêm as CVEs")
table(s, 0.6, 1.75, [2.4, 2.3, 3.0], [
    ["Imagem", "SO (C / H)", "node_modules (C / H)"],
    ["alpine", "0 / 0", "1 / 35 (app + npm da base)"],
    ["distroless", "1 / 6", "1 / 25 (só a app)"],
    ["chainguard", "0 / 0", "1 / 25 (só a app)"],
    ["slim", "1 / 56", "1 / 35"],
    ["full", "22 / 458", "1 / 35"],
], size=15, row_h=0.6)
rect(s, 8.6, 1.75, 4.15, 3.6, CARD)
text(s, 8.85, 1.95, 3.7, 0.45, "E no Go?", size=20, bold=True, color=AMBER)
text(s, 8.85, 2.5, 3.7, 2.8, [
    [("Go 1.24  ", {"bold": True, "color": CORAL}), ("23 HIGH até no scratch", {})],
    [("Go 1.27  ", {"bold": True, "color": TEAL}), ("0 CVEs", {})],
    [("A stdlib vai dentro do binário: scratch não zera CVE com toolchain velha.",
      {"color": MUTED, "size": 15})],
], size=17, space_after=12)
text(s, 0.6, 5.55, 12.1, 0.9, [[("Nas bases mínimas, ", {}),
                                ("a maioria das CVEs vem da sua app", {"bold": True, "color": AMBER}),
                                (", não da imagem base.", {})]], size=19)

# ═════════════════════════════════════════════════════════════════════════════
# 11 — Single x multi-stage
# ═════════════════════════════════════════════════════════════════════════════
s = new_slide("Rode make lab-single lab-multi primeiro. docker run handson/backend:single whoami → root. "
              "ls node_modules/.bin → tsc, jest, eslint em produção. Depois a tabela de prós e contras.")
title(s, "Single-stage x multi-stage", "Bloco 3 · 15 min")
table(s, 0.6, 1.75, [3.2, 4.45, 4.45], [
    ["", "Single-stage", "Multi-stage"],
    ["Tamanho", "toolchain + fontes + devDeps", "só o que roda"],
    ["CVEs / superfície", "alta", "baixa"],
    ["Secrets e fontes", "vazam no artefato", "ficam no stage de build"],
    ["Cache", "invalida com frequência", "stages paralelos e independentes"],
    ["Simplicidade", "✓ um FROM, fácil de ler", "mais linhas, exige entender stages"],
    ["Debug em produção", "✓ shell e ferramentas", "imagem enxuta, às vezes sem shell"],
    ["Dev / CI", "um Dockerfile só", "--target builder para testes"],
], size=15, row_h=0.56)
lab(s, "make lab-single lab-multi lab-compare")

# ═════════════════════════════════════════════════════════════════════════════
# 12 — Fluxo multi-stage
# ═════════════════════════════════════════════════════════════════════════════
s = new_slide("Ler hands-on/lab02-multi-stage/Dockerfile: deps (só produção), builder (tsc), "
              "production (copia node_modules + dist). Comparar com backend/Dockerfile e "
              "frontend/Dockerfile (Vite → nginx: node nem existe na final).")
title(s, "Três stages, uma imagem final", "Bloco 3 · lab02-multi-stage")
stages = [
    ("deps", "node:22-alpine", ["package*.json", "npm ci --omit=dev"], DOCKER),
    ("builder", "node:22-alpine", ["npm ci (com devDeps)", "COPY . .", "npm run build"], DOCKER),
    ("production", "node:22-alpine", ["COPY --from=deps node_modules", "COPY --from=builder dist",
                                      "USER node · HEALTHCHECK"], TEAL),
]
for i, (name, base, steps, color) in enumerate(stages):
    x = 0.6 + i * 4.25
    rect(s, x, 1.8, 3.75, 3.9, CARD)
    badge(s, x + 0.25, 2.0, str(i + 1), color, d=0.5)
    text(s, x + 0.9, 1.98, 2.7, 0.55, name, size=22, bold=True, anchor=MSO_ANCHOR.MIDDLE)
    text(s, x + 0.25, 2.65, 3.3, 0.4, base, size=14, color=MUTED, font=MONO)
    text(s, x + 0.25, 3.2, 3.35, 2.4, steps, size=15, font=MONO, space_after=8)
    if i < 2:
        text(s, x + 3.7, 3.35, 0.6, 0.6, "→", size=32, bold=True, color=AMBER,
             align=PP_ALIGN.CENTER)
text(s, 0.6, 5.95, 12.1, 0.5, [[("Só o stage final vai para o registry. ", {"bold": True}),
                                ("Os outros ficam no cache do build.", {"color": MUTED})]],
     size=18)
lab(s, "docker build --target builder ...")

# ═════════════════════════════════════════════════════════════════════════════
# 13 — Passo a passo
# ═════════════════════════════════════════════════════════════════════════════
s = new_slide("Passo a passo para containerizar qualquer app. Destaque: .dockerignore primeiro, "
              "ordem de camadas, exec form no CMD (PID 1 recebe SIGTERM) e secrets nunca em ARG/ENV.")
title(s, "Construindo seu Dockerfile", "Bloco 4 · passo a passo")
steps = [
    ("Entenda o runtime", "versão, porta, env, onde grava"),
    ("Escolha e fixe a base", "tag → digest"),
    (".dockerignore primeiro", "node_modules, .git, .env"),
    ("Ordem das camadas", "manifest → install → código"),
    ("Install determinístico", "npm ci · go mod download"),
    ("Multi-stage", "build separado do runtime"),
    ("Usuário não-root", "app com dono root, só leitura"),
    ("Metadados", "EXPOSE · HEALTHCHECK · LABEL"),
    ("CMD em exec form", "[\"node\", \"dist/main.js\"]"),
    ("Config fora da imagem", "secrets nunca em ARG/ENV"),
]
for i, (head, sub) in enumerate(steps):
    col, row = divmod(i, 5)
    x, y = 0.6 + col * 6.2, 1.7 + row * 0.98
    rect(s, x, y, 5.9, 0.82, CARD)
    badge(s, x + 0.15, y + 0.16, str(i + 1), DOCKER, d=0.5, size=15)
    text(s, x + 0.85, y + 0.06, 4.9, 0.4, head, size=17, bold=True)
    text(s, x + 0.85, y + 0.42, 4.9, 0.35, sub, size=14, color=MUTED)

# ═════════════════════════════════════════════════════════════════════════════
# 14 — Checklist
# ═════════════════════════════════════════════════════════════════════════════
s = new_slide("Checklist para o participante levar para os Dockerfiles do time.")
title(s, "Checklist para levar", "Bloco 4 · boas práticas")
checks = ["base oficial, versão fixa (idealmente digest)", ".dockerignore",
          "manifest de deps antes do código", "multi-stage", "USER não-root",
          "HEALTHCHECK (ou probe)", "exec form no CMD/ENTRYPOINT", "nenhum secret em camada",
          "passou no hadolint e no trivy config"]
for i, c in enumerate(checks):
    col, row = divmod(i, 5)
    x, y = 0.6 + col * 6.2, 1.8 + row * 0.92
    badge(s, x, y + 0.05, "✓", TEAL, d=0.5, size=16)
    text(s, x + 0.7, y, 5.3, 0.6, c, size=19, anchor=MSO_ANCHOR.MIDDLE)

# ═════════════════════════════════════════════════════════════════════════════
# 15 — BuildKit
# ═════════════════════════════════════════════════════════════════════════════
s = new_slide("make lab-cache: altera 1 linha de código e rebuilda. Medido: 8 s com deps antes do código "
              "x 45 s no bad-order. O cache mount ganha com rede lenta; com rede rápida não mudou o tempo. "
              "Secret mount para .npmrc. "
              "Demo de ordem de camadas: alterar comentário em main.ts e rebuildar o Dockerfile "
              "bom e o bad-order. Multi-arch só demo do palestrante.")
title(s, "BuildKit: builds rápidos e sem vazar segredo", "Bloco 4 · lab05-buildkit")
code(s, 0.6, 1.75, 7.4, 3.0, [
    "# syntax=docker/dockerfile:1",
    "RUN --mount=type=cache,target=/root/.npm \\",
    "    --mount=type=secret,id=npmrc,target=/root/.npmrc \\",
    "    npm ci",
    "",
    "docker build --secret id=npmrc,src=$HOME/.npmrc ...",
], size=14)
code(s, 0.6, 5.0, 7.4, 1.45, [
    "docker buildx build \\",
    "  --platform linux/amd64,linux/arm64 ...",
], size=14)
card(s, 8.3, 1.75, 4.45, 1.45, "ordem das camadas", ["1 linha mudada: 8 s × 45 s"], color=TEAL,
     mark="1")
card(s, 8.3, 3.37, 4.45, 1.45, "secret mount", ["token nunca vira camada"], color=AMBER, mark="2")
card(s, 8.3, 5.0, 4.45, 1.45, "cache mount", ["ganha com rede lenta"], color=DOCKER, mark="3")
lab(s, "make lab-cache")

# ═════════════════════════════════════════════════════════════════════════════
# 16 — Intervalo
# ═════════════════════════════════════════════════════════════════════════════
s = new_slide("Deixe make lab-scan rodando: o primeiro download do banco do Trivy demora.")
text(s, 0.6, 2.3, 12.1, 1.3, "Intervalo · 5 min", size=60, bold=True, align=PP_ALIGN.CENTER)
text(s, 0.6, 3.8, 12.1, 0.6, "Deixe rodando enquanto isso:", size=20, color=MUTED,
     align=PP_ALIGN.CENTER)
rect(s, 4.4, 4.6, 4.5, 0.8, CODE_BG)
text(s, 4.4, 4.6, 4.5, 0.8, "make lab-scan", size=24, bold=True, font=MONO, color=TEAL,
     align=PP_ALIGN.CENTER, anchor=MSO_ANCHOR.MIDDLE)

# ═════════════════════════════════════════════════════════════════════════════
# 17 — Scan
# ═════════════════════════════════════════════════════════════════════════════
s = new_slide("Scan não é opcional, é inventário. hadolint e trivy config antes do build; secret scan "
              "antes do commit; trivy image após o build; SBOM e assinatura na publicação; rescan "
              "contínuo em produção. Ligar com .github/workflows/appsec.yml e o DefectDojo.")
title(s, "Scan em cada etapa", "Bloco 5 · 15 min")
flow = [
    ("Escrever", "hadolint\ntrivy config", DOCKER),
    ("Commit", "trivy fs\n--scanners secret", DOCKER),
    ("Build", "trivy image\nDocker Scout · Grype", AMBER),
    ("Publicar", "SBOM\ncosign sign", TEAL),
    ("Produção", "rescan do\nregistry", TEAL),
]
for i, (stage, tools, color) in enumerate(flow):
    x = 0.6 + i * 2.5
    badge(s, x + 0.8, 1.8, str(i + 1), color, d=0.6, size=18)
    text(s, x, 2.5, 2.2, 0.45, stage, size=19, bold=True, align=PP_ALIGN.CENTER)
    rect(s, x, 3.05, 2.2, 1.3, CARD)
    text(s, x + 0.1, 3.05, 2.0, 1.3, tools.split("\n"), size=14, color=WHITE, font=MONO,
         align=PP_ALIGN.CENTER, anchor=MSO_ANCHOR.MIDDLE)
code(s, 0.6, 4.65, 12.15, 1.45, [
    "make lab-lint                      # hadolint: single x multi",
    "make lab-scan                      # trivy image + config + secret → reports/trivy-single.html",
    "trivy image --exit-code 1 --severity CRITICAL --ignore-unfixed handson/backend:multi",
], size=14)
lab(s, "make lab-lint lab-scan")

# ═════════════════════════════════════════════════════════════════════════════
# 18 — Supply chain
# ═════════════════════════════════════════════════════════════════════════════
s = new_slide("SBOM é a lista de ingredientes: dá para reescanear sem a imagem. Pin por digest + "
              "Renovate. Assinatura com cosign. Rebuild frequente: CVE corrigido na base só chega "
              "se você rebuildar.")
title(s, "Supply chain: saber o que você publica", "Bloco 5 · lab06")
sc = [
    ("SBOM", ["CycloneDX / SPDX", "trivy sbom reescaneia", "sem baixar a imagem"], DOCKER),
    ("Pin por digest", ["FROM ...@sha256", "Renovate/Dependabot", "abrem o PR de update"], DOCKER),
    ("Assinatura", ["cosign sign / verify", "política na admissão", "procedência garantida"], TEAL),
    ("Rebuild contínuo", ["base corrigida só chega", "se você rebuildar", "agende no CI"], AMBER),
]
for i, (head, lines, color) in enumerate(sc):
    card(s, 0.6 + i * 3.1, 1.8, 2.9, 3.0, head, lines, color=color, mark=str(i + 1))
text(s, 0.6, 5.15, 12.1, 0.9, [[("Gancho: ", {"bold": True, "color": AMBER}),
                                ("o secret scan NÃO acha o DD_SECRET_KEY da palestra AppSec: "
                                 "segredo genérico precisa de regra própria", {})]], size=18)
lab(s, "make lab-sbom")

# ═════════════════════════════════════════════════════════════════════════════
# 19 — Hardening
# ═════════════════════════════════════════════════════════════════════════════
s = new_slide("O frontend/Dockerfile deste repo roda nginx como root na porta 80. No lab04 vira "
              "nginx-unprivileged na 8080, rootfs read-only, cap_drop ALL, no-new-privileges, limites. "
              "Demonstrar: docker exec id, touch /etc/x falha.")
title(s, "Hardening do runtime", "Bloco 6 · 10 min")
table(s, 0.6, 1.75, [4.4, 3.5, 4.2], [
    ["Controle", "Onde", "Mitiga"],
    ["USER não-root / nginx-unprivileged", "Dockerfile", "escape com privilégio"],
    ["read_only + tmpfs", "compose / stack", "malware persistente"],
    ["cap_drop: [ALL]", "compose", "abuso de capabilities"],
    ["no-new-privileges", "compose", "escalada via setuid"],
    ["mem_limit · cpus · pids_limit", "compose", "DoS, fork bomb"],
    ["backend sem ports", "compose", "exposição desnecessária"],
], size=15, row_h=0.58)
code(s, 0.6, 5.95, 12.1, 0.62, ["docker exec handson-hardening-frontend-1 id     # uid=101, não root"],
     size=14)
lab(s, "make lab-hardening")

# ═════════════════════════════════════════════════════════════════════════════
# 20 — Comandos de rotina
# ═════════════════════════════════════════════════════════════════════════════
s = new_slide("Roteiro guiado + exercício relâmpago: quanta memória o backend usa e qual o limite? "
              "docker stats + docker inspect. Troubleshooting: exit 137, exec format error, "
              "glibc em alpine.")
title(s, "Comandos de rotina", "Bloco 7 · 10 min")
groups = [
    ("Observar", ["docker ps --format ...", "docker logs -f --tail 50", "docker stats --no-stream",
                  "docker top · docker diff"]),
    ("Investigar", ["docker exec -it <c> sh", "docker inspect -f '{{...}}'", "docker debug <c>",
                    "docker network inspect"]),
    ("Limpar", ["docker system df", "docker system prune", "docker image prune -a",
                "docker builder prune"]),
]
for i, (head, cmds) in enumerate(groups):
    x = 0.6 + i * 4.1
    rect(s, x, 1.75, 3.85, 3.05, CODE_BG)
    text(s, x + 0.25, 1.9, 3.4, 0.5, head, size=20, bold=True, color=[DOCKER, AMBER, TEAL][i])
    text(s, x + 0.25, 2.5, 3.45, 2.3, cmds, size=15, font=MONO, space_after=8)
text(s, 0.6, 5.15, 12.1, 0.9, [[("Exit 137 ", {"bold": True, "color": CORAL}), ("OOM  ·  ", {}),
                               ("exec format error ", {"bold": True, "color": CORAL}),
                               ("arquitetura  ·  ", {}),
                               ("not found ", {"bold": True, "color": CORAL}),
                               ("binário glibc em alpine", {})]], size=17)
lab(s, "hands-on/CHEATSHEET.md")

# ═════════════════════════════════════════════════════════════════════════════
# 21 — Swarm conceitos
# ═════════════════════════════════════════════════════════════════════════════
s = new_slide("Swarm = cluster de nós Docker (managers com Raft + workers). Service = estado desejado; "
              "task = container agendado. Stack = compose file. Overlay + routing mesh. Secrets "
              "nativos. Swarm x Kubernetes: simplicidade x ecossistema.")
title(s, "Docker Swarm em 1 slide", "Bloco 8 · conceitos")
rect(s, 0.6, 1.75, 7.2, 4.7, CARD)
text(s, 0.85, 1.9, 6.8, 0.45, "cluster", size=14, color=MUTED, bold=True)
nodes = [("manager", AMBER, 0.95), ("worker", DOCKER, 3.25), ("worker", DOCKER, 5.55)]
for name, color, nx in nodes:
    rect(s, nx, 2.45, 2.0, 3.1, CODE_BG)
    text(s, nx, 2.55, 2.0, 0.4, name, size=15, bold=True, color=color, align=PP_ALIGN.CENTER)
    for k in range(2):
        rect(s, nx + 0.25, 3.15 + k * 1.05, 1.5, 0.8, TEAL if k == 0 else DOCKER)
        text(s, nx + 0.25, 3.15 + k * 1.05, 1.5, 0.8, "backend" if k == 0 else "frontend",
             size=13, bold=True, color=BG, align=PP_ALIGN.CENTER, anchor=MSO_ANCHOR.MIDDLE)
text(s, 0.85, 5.75, 6.8, 0.5, "rede overlay  ·  routing mesh :8080 em qualquer nó", size=14,
     color=MUTED, align=PP_ALIGN.CENTER)
defs = [
    [("Service  ", {"bold": True, "color": DOCKER}), ("estado desejado: imagem, réplicas, limites", {})],
    [("Task  ", {"bold": True, "color": DOCKER}), ("1 container agendado num nó", {})],
    [("Stack  ", {"bold": True, "color": DOCKER}), ("services descritos num compose file", {})],
    [("Secrets  ", {"bold": True, "color": DOCKER}), ("em /run/secrets, criptografados no Raft", {})],
    [("x K8s  ", {"bold": True, "color": AMBER}), ("Swarm: simples e embutido; K8s: escala e ecossistema", {})],
]
text(s, 8.2, 1.85, 4.6, 4.6, defs, size=16, space_after=14)

# ═════════════════════════════════════════════════════════════════════════════
# 22 — Swarm hands-on
# ═════════════════════════════════════════════════════════════════════════════
s = new_slide("O rolling update troca NestJS por Go sem downtime: as duas implementações respeitam "
              "o mesmo contrato HTTP e o mesmo healthcheck. Watch em docker service ps. "
              "Multi-nó no Play with Docker: make swarm-registry + REGISTRY=127.0.0.1:5000/.")
title(s, "Rolling update: de NestJS para Go sem downtime", "Bloco 8 · lab07-swarm")
flow = [("init", "swarm-init"), ("deploy", "swarm-deploy"), ("scale", "swarm-scale"),
        ("update", "swarm-update"), ("rollback", "swarm-rollback"), ("down", "swarm-down")]
for i, (name, target) in enumerate(flow):
    x = 0.6 + i * 2.06
    rect(s, x, 1.8, 1.86, 1.45, CARD)
    badge(s, x + 0.68, 1.92, str(i + 1), TEAL if name == "update" else DOCKER, d=0.5, size=15)
    text(s, x, 2.45, 1.86, 0.4, name, size=17, bold=True, align=PP_ALIGN.CENTER)
    text(s, x, 2.82, 1.86, 0.35, target, size=12, color=MUTED, font=MONO, align=PP_ALIGN.CENTER)
code(s, 0.6, 3.55, 12.15, 2.4, [
    "docker service update --image handson/api-go:scratch bmi_backend   # 237 req, 0 falhas",
    "watch docker service ps bmi_backend        # start-first: sobe a nova, mata a antiga",
    "docker service rollback bmi_backend",
    "docker service logs bmi_backend --tail 20",
    "docker node update --availability drain <nó>    # multi-nó: tasks migram",
], size=15)
lab(s, "make swarm-init swarm-deploy swarm-update")

# ═════════════════════════════════════════════════════════════════════════════
# 22b — A jornada: quanto cada técnica enxuga
# ═════════════════════════════════════════════════════════════════════════════
JOURNEY = [  # (degrau, MB, CVEs CRITICAL+HIGH) — make lab-compare
    ("1. single-stage", MEASURED["backend:single"], 25 + 537),
    ("2. multi-stage", MEASURED["backend:alpine"], 1 + 35),
    ("3. distroless", MEASURED["backend:distroless"], 2 + 31),
    ("4. Go + scratch", MEASURED["api-go:scratch"], 0),
]


def journey_chart(s, x, y, w, h, values, number_format):
    cd = CategoryChartData()
    cd.categories = [step for step, *_ in JOURNEY]
    cd.add_series("valor", values)
    ch = s.shapes.add_chart(XL_CHART_TYPE.COLUMN_CLUSTERED, Inches(x), Inches(y),
                            Inches(w), Inches(h), cd).chart
    ch.has_legend = False
    ch.has_title = False
    ch.font.size = Pt(13)
    ch.font.name = FONT
    ch.font.color.rgb = MUTED
    plot = ch.plots[0]
    plot.gap_width = 60
    plot.has_data_labels = True
    dl = plot.data_labels
    dl.number_format = number_format
    dl.number_format_is_linked = False
    dl.position = XL_LABEL_POSITION.OUTSIDE_END
    dl.font.size = Pt(15)
    dl.font.bold = True
    dl.font.color.rgb = WHITE
    for idx in range(len(values)):
        pt = plot.series[0].points[idx]
        pt.format.fill.solid()
        pt.format.fill.fore_color.rgb = TEAL if idx == len(values) - 1 else DOCKER
    va = ch.value_axis
    va.visible = False
    va.has_major_gridlines = False
    ca = ch.category_axis
    ca.format.line.color.rgb = CARD
    ca.tick_labels.font.color.rgb = WHITE
    ca.tick_labels.font.size = Pt(13)


s = new_slide("O resumo da aula num slide: cada técnica é um degrau. Multi-stage com alpine já corta 81% "
              "do tamanho e 94% das CVEs críticas e altas; distroless tira o shell; e o binário estático em "
              "scratch chega a 9 MB e zero CVEs. Volte à votação: o que mais pesou não foi a base, foi "
              "o multi-stage e o que vai dentro da imagem.")
title(s, "Quanto cada técnica enxuga", "Encerramento · a jornada")
start_mb, start_cve = JOURNEY[0][1], JOURNEY[0][2]
end_mb = JOURNEY[-1][1]
pct = f"−{(1 - end_mb / start_mb) * 100:.1f}%".replace(".", ",")
for i, (label, value, color) in enumerate([
    ("Tamanho da imagem (MB)", pct, TEAL),
    ("CVEs CRITICAL + HIGH", f"{start_cve} → 0", TEAL),
]):
    x = 0.6 + i * 6.25
    rect(s, x, 1.65, 5.85, 4.95, CARD)
    text(s, x + 0.3, 1.8, 3.4, 0.45, label, size=17, bold=True)
    text(s, x + 3.0, 1.72, 2.6, 0.6, value, size=26, bold=True, color=color, align=PP_ALIGN.RIGHT)
    journey_chart(s, x + 0.15, 2.45, 5.55, 4.05,
                  [mb for _, mb, _ in JOURNEY] if i == 0 else [cve for *_, cve in JOURNEY],
                  "#,##0")
text(s, 0.6, 6.7, 12.1, 0.45, [[("Mesma API, mesmo contrato HTTP. ", {"bold": True}),
                                ("Só mudou a forma de construir a imagem.", {"color": MUTED})]],
     size=16)

# ═════════════════════════════════════════════════════════════════════════════
# 23 — Próximos passos
# ═════════════════════════════════════════════════════════════════════════════
s = new_slide("Recapitular o placar: de backend:single (2.070 MB, root, 25 CRITICAL e 537 HIGH) a "
              "api-go:scratch (9 MB, não-root, 0 CVEs). Próximos passos e link para a palestra AppSec.")
title(s, "Para levar para segunda-feira", "Encerramento")
nxt = [
    ("Aplicar o checklist", ["nos Dockerfiles do time", "comece pelo multi-stage"], DOCKER),
    ("Scan no pipeline", ["hadolint + trivy image", "--exit-code 1 em CRITICAL"], AMBER),
    ("Explorar", ["Docker Scout, Chainguard, DHI", "Kubernetes: probes e limits"], TEAL),
]
for i, (head, lines, color) in enumerate(nxt):
    card(s, 0.6 + i * 4.1, 1.8, 3.85, 2.4, head, lines, color=color, mark=str(i + 1))
text(s, 0.6, 4.6, 12.1, 1.4, [
    [("Repositório: ", {"bold": True}), ("hands-on/README.md  ·  roteiro-containerizacao.md",
                                          {"font": MONO, "color": MUTED, "size": 16})],
    [("Palestra AppSec: ", {"bold": True}), ("roteiro-palestra.md  (DAST, SCA, DefectDojo)",
                                              {"font": MONO, "color": MUTED, "size": 16})],
], size=18, space_after=10)

# ═════════════════════════════════════════════════════════════════════════════
# 24 — Perguntas
# ═════════════════════════════════════════════════════════════════════════════
s = new_slide("Puxar o debate: Alpine ou distroless para Node em produção? Quem já teve problema "
              "de módulo nativo em Alpine? Vale Swarm em 2026?")
text(s, 0.6, 1.9, 12.1, 1.3, "Perguntas?", size=66, bold=True, align=PP_ALIGN.CENTER)
qs = ["Revotação: Alpine ou distroless?",
      "Quem já teve módulo nativo quebrando no Alpine?",
      "Vale a pena Swarm em 2026?"]
for i, q in enumerate(qs):
    x = 0.6 + i * 4.1
    rect(s, x, 3.9, 3.85, 1.5, CARD)
    text(s, x + 0.25, 3.9, 3.35, 1.5, q, size=17, color=MUTED, anchor=MSO_ANCHOR.MIDDLE,
         align=PP_ALIGN.CENTER)


# ── Salvar ───────────────────────────────────────────────────────────────────
default_out = os.path.join(os.path.dirname(__file__), "..", "..", "reports", "containerizacao-palestra.pptx")
out = os.path.abspath(os.environ.get("PPTX_OUT", default_out))   # PPTX_OUT: gerar cópia com o deck aberto
prs.save(out)
print(f"✔ {len(prs.slides)} slides → {out}")
